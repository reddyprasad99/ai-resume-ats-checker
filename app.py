import io
import re
from collections import Counter

import streamlit as st
import fitz  # PyMuPDF
from docx import Document

st.set_page_config(page_title="AI Resume & ATS Checker", page_icon="📄", layout="wide")

st.title("📄 Resume & ATS Checker")
st.caption("Check how closely your resume matches a job description. Your files are processed for this session only; this app does not intentionally save them.")

st.info(
    "Privacy note: This app does not save uploaded resumes to a database or write them to disk. "
    "The hosting provider may still process technical logs or temporary memory, and downloaded files are saved by your browser."
)

STOPWORDS = {
    "a", "an", "and", "are", "as", "at", "be", "been", "by", "for", "from", "has",
    "have", "in", "is", "it", "of", "on", "or", "our", "that", "the", "their",
    "this", "to", "was", "were", "will", "with", "you", "your", "we", "they",
    "using", "use", "work", "working", "role", "job", "candidate", "ability",
    "skills", "skill", "experience", "required", "preferred", "including"
}

def extract_text(uploaded_file):
    """Extract text from an uploaded PDF or DOCX held in memory."""
    data = uploaded_file.getvalue()
    name = uploaded_file.name.lower()

    if name.endswith(".pdf"):
        chunks = []
        with fitz.open(stream=data, filetype="pdf") as pdf:
            for page in pdf:
                chunks.append(page.get_text())
        return "\n".join(chunks)

    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        chunks = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                chunks.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(chunks)

    raise ValueError("Please upload a PDF or DOCX file.")

def words(text):
    found = re.findall(r"[A-Za-z][A-Za-z0-9+#.-]{1,}", text.lower())
    return [w.strip(".-") for w in found if len(w.strip(".-")) > 1 and w not in STOPWORDS]

def keyword_analysis(resume_text, job_text):
    resume_words = set(words(resume_text))
    job_counts = Counter(words(job_text))
    job_keywords = [word for word, _ in job_counts.most_common(40)]
    matched = [word for word in job_keywords if word in resume_words]
    missing = [word for word in job_keywords if word not in resume_words]
    score = round(100 * len(matched) / max(len(job_keywords), 1))
    return score, matched, missing, job_keywords

def build_report_text(score, matched, missing, suggestions):
    lines = [
        "RESUME & ATS CHECK REPORT",
        "=" * 28,
        f"Estimated keyword match: {score}/100",
        "",
        "Matched terms:",
        ", ".join(matched) if matched else "No clear matches detected.",
        "",
        "Potentially missing terms (verify they truthfully describe your skills):",
        ", ".join(missing[:25]) if missing else "No major missing terms detected.",
        "",
        "Suggestions:",
        *[f"- {item}" for item in suggestions],
        "",
        "Note: This is an estimated keyword comparison, not an actual score from an employer's ATS."
    ]
    return "\n".join(lines)

resume_file = st.file_uploader("Upload your resume (PDF or DOCX)", type=["pdf", "docx"])
job_text = st.text_area("Paste the job description", height=220, placeholder="Paste the complete job description here...")

if st.button("Analyze resume", type="primary", use_container_width=True):
    if not resume_file:
        st.error("Upload a PDF or DOCX resume first.")
    elif not job_text.strip():
        st.error("Paste the job description first.")
    else:
        try:
            resume_text = extract_text(resume_file)
            if not resume_text.strip():
                st.error("No readable text was found. If this is a scanned PDF, run OCR on your own device and upload a text-based PDF.")
            else:
                score, matched, missing, job_keywords = keyword_analysis(resume_text, job_text)
                suggestions = []
                if len(resume_text.split()) < 150:
                    suggestions.append("Your resume appears short. Add relevant education, projects, skills, or measurable achievements where accurate.")
                if not re.search(r"\b(email|@)\b", resume_text, re.I) and "@" not in resume_text:
                    suggestions.append("Check that your contact email is clearly visible.")
                if not re.search(r"\b(education|degree|bachelor|university|college)\b", resume_text, re.I):
                    suggestions.append("Consider adding a clearly labelled Education section if relevant.")
                if not re.search(r"\b(project|projects)\b", resume_text, re.I):
                    suggestions.append("If relevant, add a Projects section with your contribution and tools used.")
                if missing:
                    suggestions.append("Review the missing terms below and add only those that truthfully match your qualifications.")
                suggestions.append("Use clear section headings, consistent dates, and simple formatting; avoid tables, graphics, and text boxes when possible.")
                suggestions.append("Check spelling, contact details, dates, and every claim before applying.")

                st.subheader("Results")
                c1, c2, c3 = st.columns(3)
                c1.metric("Estimated keyword match", f"{score}/100")
                c2.metric("Matched terms", len(matched))
                c3.metric("Potential missing terms", len(missing))
                st.progress(score / 100, text="Keyword overlap only — not a guaranteed ATS pass score.")

                left, right = st.columns(2)
                with left:
                    st.markdown("#### Matched keywords")
                    st.write(", ".join(matched) if matched else "No matching keywords detected.")
                with right:
                    st.markdown("#### Potentially missing keywords")
                    st.write(", ".join(missing[:25]) if missing else "No major gaps detected.")

                st.markdown("#### Improvement suggestions")
                for item in suggestions:
                    st.write(f"- {item}")

                report = build_report_text(score, matched, missing, suggestions)
                st.download_button(
                    "Download analysis report (.txt)",
                    data=report.encode("utf-8"),
                    file_name="resume_ats_report.txt",
                    mime="text/plain",
                    use_container_width=True
                )
                st.caption("This initial version does not automatically rewrite or create a new resume. It produces a downloadable analysis report.")
                # Avoid retaining extracted text in session state; do not assign it to st.session_state.
                del resume_text
        except Exception:
            st.error("We couldn't read this file. Try a valid text-based PDF or DOCX. No document content is shown in the error message.")

st.divider()
st.caption("Privacy-first prototype • No login • No resume database • No external AI API in this version")
